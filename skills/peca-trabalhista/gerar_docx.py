#!/usr/bin/env python3
"""Gera .docx de peça trabalhista aplicando template visual do Ferreira & Martins.

Recebe um arquivo Markdown estruturado e produz um .docx pronto para protocolo,
herdando do template `peca_modelo_FM.docx` o cabeçalho (logo), o rodapé
(contatos das filiais), as margens A4 e os estilos.

Convenções do Markdown de entrada:

  # TÍTULO DA PEÇA           -> centralizado, negrito, MAIÚSCULAS, 14pt
  ## I — DOS FATOS           -> negrito, esquerda, 12pt (cabeçalho de seção)
  ### I.1 — Subseção         -> negrito itálico, esquerda, 12pt
  **texto em negrito**       -> negrito inline
  *texto em itálico*         -> itálico inline
  > citação                  -> recuo esquerdo 2,5 cm, justificado (lei/jurisprudência)
  ^ texto                    -> centralizado (fecho, local/data, assinatura)
  ---                        -> parágrafo em branco
  qualquer outra linha       -> justificada, 12pt, sem recuo

Uso:
  python3 gerar_docx.py --corpo-arquivo peca.md --saida ~/Downloads/saida.docx
  python3 gerar_docx.py --corpo-arquivo peca.md --saida ~/Downloads/saida.docx --titulo "CONTESTACAO"
"""

import argparse
import copy
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


TEMPLATE = Path(__file__).parent / "templates" / "peca_modelo_FM.docx"
FONT_NAME = "Times New Roman"
FONT_SIZE = Pt(12)
TITLE_SIZE = Pt(14)


def clean_body(doc: Document) -> None:
    """Remove todos os parágrafos do corpo, preservando seções/header/footer."""
    body = doc.element.body
    # Manter os <w:sectPr> que ficam ao final do body (definem layout/margens/header/footer)
    children = list(body)
    for child in children:
        if child.tag == qn("w:p") or child.tag == qn("w:tbl"):
            body.remove(child)


def add_paragraph(doc: Document, text: str, alignment, *, bold: bool = False,
                  italic: bool = False, size: Pt = FONT_SIZE,
                  inline_format: bool = True, left_indent=None):
    """Adiciona parágrafo com formatação. Suporta **bold** e *italic* inline."""
    p = doc.add_paragraph()
    p.alignment = alignment
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.left_indent = left_indent if left_indent is not None else Pt(0)
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    if not inline_format:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = FONT_NAME
        run.font.size = size
        return p

    # Parser inline: **bold** e *italic*
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*"):
            run = p.add_run(part[1:-1])
            run.italic = True
        else:
            run = p.add_run(part)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        run.font.name = FONT_NAME
        run.font.size = size
    return p


def render_markdown(doc: Document, md_text: str) -> None:
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if not line:
            add_paragraph(doc, "", WD_ALIGN_PARAGRAPH.JUSTIFY)
            i += 1
            continue

        if line.strip() == "---":
            add_paragraph(doc, "", WD_ALIGN_PARAGRAPH.JUSTIFY)
            i += 1
            continue

        # Heading 1: título da peça (centralizado, negrito, MAIÚSCULAS, 14pt)
        if line.startswith("# ") and not line.startswith("## "):
            text = line[2:].strip().upper()
            add_paragraph(
                doc, text, WD_ALIGN_PARAGRAPH.CENTER,
                bold=True, size=TITLE_SIZE, inline_format=False,
            )
            i += 1
            continue

        # Heading 3: subseção (negrito itálico, esquerda)
        if line.startswith("### "):
            text = line[4:].strip()
            add_paragraph(
                doc, text, WD_ALIGN_PARAGRAPH.LEFT,
                bold=True, italic=True, inline_format=False,
            )
            i += 1
            continue

        # Heading 2: seção (negrito, esquerda)
        if line.startswith("## "):
            text = line[3:].strip()
            add_paragraph(
                doc, text, WD_ALIGN_PARAGRAPH.LEFT,
                bold=True, inline_format=False,
            )
            i += 1
            continue

        # Bloco citado (>): citação de lei/jurisprudência — recuo esquerdo 2,5cm, justificado
        if line.startswith(">"):
            text = line.lstrip(">").strip()
            add_paragraph(
                doc, text, WD_ALIGN_PARAGRAPH.JUSTIFY,
                left_indent=Cm(2.5),
            )
            i += 1
            continue

        # Centralizado (^): fecho, data/local, assinatura
        if line.startswith("^ ") or line == "^":
            text = line[1:].strip()
            add_paragraph(doc, text, WD_ALIGN_PARAGRAPH.CENTER)
            i += 1
            continue

        # Linha que abre com "AO JUÍZO" ou "AO EGRÉGIO" -> negrito MAIÚSCULAS justificado
        if re.match(r"^\s*(AO|AOS|EXC|EGRÉGIO|AO EGRÉGIO|EXCELENTÍSSIMO)", line, re.IGNORECASE) and line.isupper():
            add_paragraph(doc, line, WD_ALIGN_PARAGRAPH.JUSTIFY,
                          bold=True, inline_format=False)
            i += 1
            continue

        # Parágrafo justificado padrão (com inline format)
        add_paragraph(doc, line, WD_ALIGN_PARAGRAPH.JUSTIFY)
        i += 1


def main():
    parser = argparse.ArgumentParser(
        description="Gera .docx de peça trabalhista aplicando template FM.",
    )
    parser.add_argument("--corpo-arquivo", required=True,
                        help="Arquivo .md com o conteúdo da peça.")
    parser.add_argument("--saida", required=True,
                        help="Caminho do .docx de saída.")
    parser.add_argument("--titulo", default=None,
                        help="Título opcional adicionado ao início (centralizado, MAIÚSCULAS).")
    parser.add_argument("--template", default=str(TEMPLATE),
                        help=f"Template base (padrão: {TEMPLATE}).")
    args = parser.parse_args()

    md_path = Path(args.corpo_arquivo).expanduser()
    if not md_path.exists():
        sys.exit(f"Erro: arquivo {md_path} não encontrado.")

    template_path = Path(args.template).expanduser()
    if not template_path.exists():
        sys.exit(f"Erro: template {template_path} não encontrado.")

    saida_path = Path(args.saida).expanduser()
    saida_path.parent.mkdir(parents=True, exist_ok=True)

    md_text = md_path.read_text(encoding="utf-8")

    doc = Document(str(template_path))
    clean_body(doc)

    if args.titulo:
        add_paragraph(
            doc, args.titulo.upper(), WD_ALIGN_PARAGRAPH.CENTER,
            bold=True, size=TITLE_SIZE, inline_format=False,
        )
        add_paragraph(doc, "", WD_ALIGN_PARAGRAPH.JUSTIFY)

    render_markdown(doc, md_text)

    doc.save(str(saida_path))
    print(f"OK -> {saida_path}")


if __name__ == "__main__":
    main()
